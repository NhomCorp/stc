"""Tạo data/log_export.json từ data/sheet_export.xlsx (Log_MM_YYYY)."""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "sheet_export.xlsx"
OUT = ROOT / "data" / "log_export.json"


def ser(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"Thiếu {XLSX}")

    wb = load_workbook(XLSX, data_only=True, read_only=True)
    all_rows = []
    for name in wb.sheetnames:
        if not (name.startswith("Log_") and name != "Log_Chuyen"):
            continue
        ws = wb[name]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < 2:
                continue
            vals = list(row[:9]) if row else []
            if not any(v is not None and str(v).strip() != "" for v in vals):
                continue
            all_rows.append(
                {
                    "sourceSheet": name,
                    "sourceRowIndex": i + 1,
                    "ngay": ser(vals[0]),
                    "phanLoai": None if vals[1] is None else str(vals[1]).strip(),
                    "soTien": vals[2],
                    "vi": None if vals[3] is None else str(vals[3]).strip(),
                    "doiTuong": None if vals[4] is None else str(vals[4]).strip(),
                    "danhMuc": None if vals[5] is None else str(vals[5]).strip(),
                    "ghiChu": None if vals[6] is None else str(vals[6]).strip(),
                    "uniqueKey": None if vals[7] is None else str(vals[7]).strip(),
                    "status": None if vals[8] is None else str(vals[8]).strip(),
                }
            )
    wb.close()
    OUT.write_text(json.dumps(all_rows, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"Wrote {len(all_rows)} rows -> {OUT}")


if __name__ == "__main__":
    main()
